"""
AI 回覆每日審核報告：把某一天所有 AI 生成的內容匯出成 Excel，供研究者逐則驗核。

兩個工作表：
- 「作答解說」：作答後的 AI 解說（feedback_logs），附題目、選項、正解、學生選的、對錯
- 「AI助教」：AI 助教的每一輪問答（ai_conversation_log），同一次對話依時間排在一起

每列最後有「審核」「備註」兩欄留白，給研究者填「正確／有誤」與錯誤說明；審核完的
Excel 本身就是人工檢核紀錄。使用者以開通碼標示（不出現 LINE ID），開通碼對應到誰
請看 access_codes/ 底下的 CSV。輸出到 ai_reviews/（已列入 .gitignore，內含學生作答）。

日期以台灣時間（UTC+8）切分，預設是昨天。

用法：
    python scripts/ai_review_report.py                  # 昨天
    python scripts/ai_review_report.py --date 2026-10-20
    python scripts/ai_review_report.py --date 2026-10-20 --days 7   # 從該日起連續 7 天
"""

import argparse
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.db.client import supabase  # noqa: E402
from app.services.flex_templates import MODE_LABELS  # noqa: E402
from app.services.question_picker import option_text  # noqa: E402

TW = timezone(timedelta(hours=8))
OUT_DIR = Path(__file__).resolve().parent.parent / "ai_reviews"
STAGE_LABELS = {"semantic_choice": "意思題", "situational_choice": "情境題"}
KIND_LABELS = {"initial": "初次解析", "followup": "追問", "lookup_failed": "題號查無", "limit_reached": "額度用完"}


def _fetch_all(query_fn) -> list[dict]:
    rows, start = [], 0
    while True:
        batch = query_fn().range(start, start + 999).execute().data
        rows += batch
        if len(batch) < 1000:
            return rows
        start += 1000


def _fetch_by_ids(table: str, ids: set, cols: str = "*", key: str = "id") -> dict:
    ids = [i for i in ids if i]
    out = {}
    for i in range(0, len(ids), 200):
        for row in supabase.table(table).select(cols).in_(key, ids[i : i + 200]).execute().data:
            out[row[key]] = row
    return out


def _labeled(question: dict, ids: str) -> str:
    return "、".join(f"{i.upper()}. {option_text(question, i)}" for i in (ids or "").split("、") if i)


def _to_tw(ts: str) -> str:
    # Supabase 的小數秒位數不固定（例如 .30833），Python 3.9 的 fromisoformat 不接受，先去掉
    ts = re.sub(r"\.\d+", "", ts.replace("Z", "+00:00"))
    return datetime.fromisoformat(ts).astimezone(TW).strftime("%m/%d %H:%M:%S")


def _user_codes(user_ids: set) -> dict:
    codes = _fetch_by_ids("access_codes", user_ids, "code, category, user_id", key="user_id")
    return {uid: (row["code"], "測試" if row["category"] == "tester" else "實驗組") for uid, row in codes.items()}


def _write_sheet(ws, headers: list[str], widths: list[int], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E4DCC8")
    for row in rows:
        ws.append(row)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def build_report(start: date, days: int) -> Path:
    since = datetime(start.year, start.month, start.day, tzinfo=TW)
    until = since + timedelta(days=days)
    s, u = since.isoformat(), until.isoformat()

    feedback = _fetch_all(
        lambda: supabase.table("feedback_logs")
        .select("id, attempt_log_id, ai_generated_text, created_at")
        .gte("created_at", s)
        .lt("created_at", u)
        .order("created_at")
    )
    tutor = _fetch_all(
        lambda: supabase.table("ai_conversation_log")
        .select("user_id, question_id, role, message, kind, conversation_id, created_at")
        .gte("created_at", s)
        .lt("created_at", u)
        .order("created_at")
    )

    attempts = _fetch_by_ids("attempts_log", {f["attempt_log_id"] for f in feedback})
    questions = _fetch_by_ids(
        "questions",
        {a["question_id"] for a in attempts.values()} | {t["question_id"] for t in tutor},
    )
    codes = _user_codes({a["user_id"] for a in attempts.values()} | {t["user_id"] for t in tutor})

    feedback_rows = []
    for f in feedback:
        a = attempts.get(f["attempt_log_id"]) or {}
        q = questions.get(a.get("question_id")) or {}
        code, group = codes.get(a.get("user_id"), ("", ""))
        source = "每日挑戰" if a.get("daily_challenge_id") else ("錯題複習" if a.get("attempt_type") == "review" else "練習")
        feedback_rows.append(
            [
                _to_tw(f["created_at"]),
                code,
                group,
                MODE_LABELS.get(q.get("mode"), q.get("mode")),
                q.get("question_number"),
                STAGE_LABELS.get(q.get("stage"), ""),
                source,
                q.get("context_sentence") or "",
                "\n".join(f"{o['id'].upper()}. {o['text']}" for o in q.get("options") or []),
                _labeled(q, q.get("correct_option")),
                _labeled(q, a.get("selected_option")),
                "對" if a.get("is_correct") else "錯",
                f["ai_generated_text"],
                "",
                "",
            ]
        )

    # AI 助教：同一次對話排在一起，再依對話開始時間排序
    first_seen: dict = {}
    for t in tutor:
        first_seen.setdefault(t["conversation_id"] or t["created_at"], t["created_at"])
    tutor.sort(key=lambda t: (first_seen[t["conversation_id"] or t["created_at"]], t["created_at"]))

    tutor_rows = []
    pending_user: dict = {}
    for t in tutor:
        if t["role"] == "user":
            if t["kind"] in ("lookup_failed", "limit_reached"):
                q = questions.get(t["question_id"]) or {}
                code, group = codes.get(t["user_id"], ("", ""))
                tutor_rows.append(
                    [_to_tw(t["created_at"]), code, group, MODE_LABELS.get(q.get("mode"), ""), q.get("question_number"),
                     KIND_LABELS.get(t["kind"], t["kind"]), q.get("context_sentence") or "", t["message"], "（無 AI 回覆）", "", ""]
                )
            else:
                pending_user[t["conversation_id"]] = t
            continue
        asked = pending_user.pop(t["conversation_id"], None)
        q = questions.get(t["question_id"]) or {}
        code, group = codes.get(t["user_id"], ("", ""))
        tutor_rows.append(
            [
                _to_tw(t["created_at"]),
                code,
                group,
                MODE_LABELS.get(q.get("mode"), ""),
                q.get("question_number"),
                KIND_LABELS.get(t["kind"], t["kind"] or ""),
                q.get("context_sentence") or "",
                asked["message"] if asked else "",
                t["message"],
                "",
                "",
            ]
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "作答解說"
    _write_sheet(
        ws,
        ["時間", "開通碼", "組別", "模式", "題號", "題型", "來源", "題目", "選項", "正解", "學生選的", "對錯",
         "AI 解說", "審核（正確／有誤）", "備註"],
        [14, 9, 8, 8, 6, 8, 9, 30, 22, 16, 16, 5, 60, 12, 30],
        feedback_rows,
    )
    _write_sheet(
        wb.create_sheet("AI助教"),
        ["時間", "開通碼", "組別", "模式", "題號", "類型", "題目", "學生訊息", "AI 回覆", "審核（正確／有誤）", "備註"],
        [14, 9, 8, 8, 6, 9, 30, 25, 60, 12, 30],
        tutor_rows,
    )

    OUT_DIR.mkdir(exist_ok=True)
    label = start.isoformat() if days == 1 else f"{start.isoformat()}_{days}天"
    out_path = OUT_DIR / f"AI回覆審核_{label}.xlsx"
    wb.save(out_path)
    print(f"作答解說 {len(feedback_rows)} 則、AI 助教 {len(tutor_rows)} 則 → {out_path}")
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description="匯出 AI 回覆審核報告（Excel）")
    parser.add_argument("--date", help="起始日期 YYYY-MM-DD（台灣時間），預設昨天")
    parser.add_argument("--days", type=int, default=1, help="涵蓋天數，預設 1")
    args = parser.parse_args()
    start = date.fromisoformat(args.date) if args.date else datetime.now(TW).date() - timedelta(days=1)
    build_report(start, args.days)


if __name__ == "__main__":
    main()
